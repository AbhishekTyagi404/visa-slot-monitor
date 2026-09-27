import sys
import time
import os
import csv
import io
import smtplib
from email.mime.text import MIMEText
from selenium import webdriver
from selenium.webdriver.firefox.options import Options
from selenium.webdriver.firefox.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, WebDriverException
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from datetime import datetime

# ==== CONFIG ====
# Pick the visa to monitor: run `python visa_slot_monitor.py F1` or `python visa_slot_monitor.py J1`
# (defaults to DEFAULT_VISA if no argument is given).
# The site now has a separate page per visa type (the old #F1Regular anchor is gone),
# and the heading changed from "Last F-1 (Regular) Availability" to "Latest F-1 (Regular) Availability".
VISA_TYPES = {
    "F1": {"slug": "f-1-regular", "label": "F-1 (Regular)", "filter": "F-1"},
    "J1": {"slug": "j-1-regular", "label": "J-1 (Regular)", "filter": "J-1"},
    # Add more here, e.g. "H1B": {"slug": "h-1b-regular", "label": "H-1B (Regular)", "filter": "H-1B"},
}
DEFAULT_VISA = "F1"

VISA_KEY = (sys.argv[1] if len(sys.argv) > 1 else DEFAULT_VISA).upper().replace("-", "")
if VISA_KEY not in VISA_TYPES:
    sys.exit(f"Unknown visa '{VISA_KEY}'. Choose one of: {', '.join(VISA_TYPES)}")
VISA = VISA_TYPES[VISA_KEY]

URL = f"https://checkvisaslots.com/latest-us-visa-availability/{VISA['slug']}/"
TABLE_HEADING = f"Latest {VISA['label']} Availability"
VISA_TYPE_FILTER = VISA["filter"]   # only keep rows whose Visa Type contains this

EXCEL_PATH = f"{VISA_KEY.lower()}_visa_slot_log.xlsx"
TEXT_LOG = f"{VISA_KEY.lower()}_text.txt"   # one state file per visa type
CHECK_INTERVAL = 60               # seconds between refreshes (site updates every ~2 min)

SENDER_EMAIL = "sender-email@gmail.com"
SENDER_PASSWORD = "your-app-password"
RECEIVER_EMAILS = [
    "reciever-email1@gmail.com",
    "reciever-email2@gmail.com"
]

# Current table columns: Visa Location | Visa Type | Earliest Date | Slots on Earliest Date |
#                        Total Dates Available | Last Seen At | Relative Time
# ("Relative Time" is left out on purpose - it changes constantly and would cause false alerts)
COLUMNS_TO_KEEP = ['visa location', 'visa type', 'earliest date', 'slots on earliest date', 'last seen at']


# ==== EMAIL ALERT FUNCTION ====
def send_email_alert(changed_rows):
    first_row = changed_rows[0]
    # first_row = [location, visa type, earliest date, slots, last seen at]
    subject = f"{VISA['label']} Slot Changed: {first_row[0]} - earliest {first_row[2]} (seen {first_row[4]})"

    body_lines = [f"Changes detected in {VISA['label']} visa slots:\n",
                  " | ".join(col.title() for col in COLUMNS_TO_KEEP)]
    for row in changed_rows:
        body_lines.append(" | ".join(row))
    body_lines.append(f"\nSource: {URL}")
    body = "\n".join(body_lines)

    msg = MIMEText(body)
    msg['From'] = SENDER_EMAIL
    msg['To'] = ", ".join(RECEIVER_EMAILS)
    msg['Subject'] = subject

    try:
        server = smtplib.SMTP('smtp.gmail.com', 587)
        server.starttls()
        server.login(SENDER_EMAIL, SENDER_PASSWORD)
        server.sendmail(SENDER_EMAIL, RECEIVER_EMAILS, msg.as_string())
        server.quit()
        print("📧 Email alert sent.")
    except Exception as e:
        print(f"❌ Failed to send email: {e}")


# ==== EXCEL FUNCTIONS ====
def initialize_excel():
    if not os.path.exists(EXCEL_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = f"{VISA_KEY}VisaSlots"
        ws.append(["Timestamp"] + [col.title() for col in COLUMNS_TO_KEEP])
        wb.save(EXCEL_PATH)

def append_to_excel(rows):
    wb = load_workbook(EXCEL_PATH)
    ws = wb.active
    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    for row in rows:
        ws.append([timestamp] + row)
    wb.save(EXCEL_PATH)


# ==== TEXT LOG FUNCTIONS ====
def read_old_text():
    if os.path.exists(TEXT_LOG):
        with open(TEXT_LOG, 'r', encoding='utf-8', newline='') as f:
            return f.read()
    return ""

def save_text(text):
    with open(TEXT_LOG, 'w', encoding='utf-8', newline='') as f:
        f.write(text)

def get_rows_from_csv(csv_text):
    reader = csv.reader(io.StringIO(csv_text))
    next(reader, None)  # Skip header
    return [row for row in reader]


# ==== TABLE PARSING ====
def normalize(text):
    return " ".join(text.split()).lower()

def get_visa_table(driver):
    soup = BeautifulSoup(driver.page_source, 'html.parser')

    # 1) Try: table that follows the "Latest <visa> Availability" heading.
    #    find_next (not find_next_sibling) because the table may be wrapped in a div.
    header = soup.find(lambda tag: tag.name in ["h1", "h2", "h3", "h4", "strong"]
                       and TABLE_HEADING.lower() in normalize(tag.get_text()))
    if header:
        table = header.find_next('table')
        if table:
            return table

    # 2) Fallback: any table whose header row contains "visa location"
    for table in soup.find_all('table'):
        first_row = table.find('tr')
        if first_row and 'visa location' in normalize(first_row.get_text(" ")):
            return table

    print(f"❌ {VISA['label']} availability table not found.")
    return None

def extract_relevant_columns(table_tag):
    rows = table_tag.find_all('tr')
    if not rows:
        return [], ""

    headers = [normalize(th.get_text(" ")) for th in rows[0].find_all(['th', 'td'])]
    indexes = []
    for col in COLUMNS_TO_KEEP:
        if col in headers:
            indexes.append(headers.index(col))
        else:
            print(f"❌ Column '{col}' not found. Columns on page: {headers}")
            return [], ""

    type_idx = headers.index('visa type')

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([col.title() for col in COLUMNS_TO_KEEP])

    extracted_rows = []
    for row in rows[1:]:
        cells = row.find_all(['td', 'th'])
        if len(cells) < max(indexes) + 1:
            continue
        if VISA_TYPE_FILTER.lower() not in cells[type_idx].get_text(strip=True).lower():
            continue
        selected_cells = [" ".join(cells[i].get_text(" ", strip=True).split()) for i in indexes]
        extracted_rows.append(selected_cells)
        writer.writerow(selected_cells)

    return extracted_rows, output.getvalue()


def load_page(driver, first=False):
    """Load/refresh the page and wait until a table is present."""
    if first:
        driver.get(URL)
    else:
        driver.refresh()
    WebDriverWait(driver, 20).until(EC.presence_of_element_located((By.TAG_NAME, "table")))


# ==== MAIN LOOP ====
def main():
    options = Options()
    # options.add_argument('--headless')  # Uncomment to run in background

    service = Service()
    driver = webdriver.Firefox(service=service, options=options)

    print(f"👀 Monitoring {VISA['label']} at {URL}")
    try:
        initialize_excel()
        old_text = read_old_text()
        old_rows = get_rows_from_csv(old_text) if old_text else []
        first = True

        while True:
            try:
                load_page(driver, first=first)
                first = False
            except (TimeoutException, WebDriverException) as e:
                print(f"⚠️ Page load problem ({type(e).__name__}), retrying in {CHECK_INTERVAL}s...")
                time.sleep(CHECK_INTERVAL)
                continue

            table_tag = get_visa_table(driver)
            if table_tag is not None:
                new_rows, new_csv = extract_relevant_columns(table_tag)

                if new_csv and new_csv != old_text:
                    print(f"⚠️ Change detected! [{datetime.now():%H:%M:%S}]")

                    changed_rows = [row for row in new_rows if row not in old_rows]
                    if changed_rows:
                        print("\n🆕 New/updated rows:")
                        for row in changed_rows:
                            print(" → ", row)
                        send_email_alert(changed_rows)
                        append_to_excel(changed_rows)
                    else:
                        print("⚠️ Change detected, but no new rows (a row was probably removed).")

                    save_text(new_csv)
                    old_text = new_csv
                    old_rows = new_rows
                elif new_csv:
                    print(f"✅ No change detected. [{datetime.now():%H:%M:%S}]")

            time.sleep(CHECK_INTERVAL)

    except KeyboardInterrupt:
        print("\n⛔ Script stopped by user.")
    finally:
        driver.quit()

if __name__ == "__main__":
    main()
